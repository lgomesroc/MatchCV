from io import BytesIO

import pytest
from docx import Document

from MatchCV.Domain.Enums.FileType import FileType
from MatchCV.Parser.Exceptions.ParserException import ParserException
from MatchCV.Parser.Parsers.DocxResumeParser import DocxResumeParser


def create_docx_bytes():
    document = Document()
    document.add_paragraph(
        "Luciano Rocha - Desenvolvedor de software"
    )
    document.add_paragraph(
        "Experiência com desenvolvimento e testes automatizados."
    )

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)

    return stream.getvalue()


class TestDocxResumeParser:
    def setup_method(self):
        self.parser = DocxResumeParser()

    def test_extracts_paragraphs_from_docx(self):
        file_bytes = create_docx_bytes()

        result = self.parser.parse(
            file_stream=BytesIO(file_bytes),
            file_name="curriculo.docx",
            file_size_bytes=len(file_bytes),
        )

        assert result.file_name == "curriculo.docx"
        assert result.file_type == FileType.DOCX
        assert result.file_size_bytes == len(file_bytes)
        assert "Luciano Rocha" in result.extracted_text
        assert "testes automatizados" in result.extracted_text
        assert result.is_text_extractable is True
        assert result.is_single_column is True
        assert result.has_images is False
        assert result.is_password_protected is False

    def test_extracts_text_from_tables(self):
        document = Document()
        table = document.add_table(
            rows=1,
            cols=2,
        )
        table.cell(0, 0).text = "Tecnologia"
        table.cell(0, 1).text = "Python"

        stream = BytesIO()
        document.save(stream)
        file_bytes = stream.getvalue()

        result = self.parser.parse(
            file_stream=BytesIO(file_bytes),
            file_name="curriculo.docx",
            file_size_bytes=len(file_bytes),
        )

        assert "Tecnologia | Python" in result.extracted_text

    def test_marks_text_as_not_extractable_when_document_is_empty(self):
        document = Document()
        stream = BytesIO()
        document.save(stream)
        file_bytes = stream.getvalue()

        result = self.parser.parse(
            file_stream=BytesIO(file_bytes),
            file_name="vazio.docx",
            file_size_bytes=len(file_bytes),
        )

        assert result.extracted_text == ""
        assert result.is_text_extractable is False
        assert result.page_count == 1

    def test_raises_when_docx_cannot_be_opened(self):
        with pytest.raises(
            ParserException,
            match="Não foi possível ler o arquivo DOCX",
        ):
            self.parser.parse(
                file_stream=BytesIO(b"not a valid docx"),
                file_name="invalido.docx",
                file_size_bytes=16,
            )

    def test_rejects_unsupported_extension(self):
        with pytest.raises(
            ParserException,
            match="Formato de currículo não suportado",
        ):
            self.parser.parse(
                file_stream=BytesIO(b"content"),
                file_name="curriculo.txt",
                file_size_bytes=7,
            )

    def test_rejects_file_larger_than_one_mb(self):
        with pytest.raises(
            ParserException,
            match="ultrapassar 1 MB",
        ):
            self.parser.parse(
                file_stream=BytesIO(b"content"),
                file_name="curriculo.docx",
                file_size_bytes=1_048_577,
            )

    def test_estimates_one_page_for_empty_text(self):
        assert DocxResumeParser._estimate_page_count("") == 1

    def test_estimates_pages_from_explicit_page_breaks(self):
        assert (
            DocxResumeParser._estimate_page_count(
                "Página 1\fPágina 2"
            )
            == 2
        )

    def test_estimates_at_least_one_page(self):
        assert (
            DocxResumeParser._estimate_page_count("Texto")
            == 1
        )

    def test_detects_normal_document_as_single_column(self):
        document = Document()
        document.add_paragraph("Texto de exemplo")

        assert (
            DocxResumeParser._detect_single_column(document)
            is True
        )

    def test_returns_none_for_document_without_paragraphs_or_tables(self):
        document = Document()

        assert (
            DocxResumeParser._detect_single_column(document)
            is None
        )
