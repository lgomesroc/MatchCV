from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter

from MatchCV.Parser.Exceptions.ParserException import ParserException
from MatchCV.Parser.Parsers.DocResumeParser import DocResumeParser
from MatchCV.Parser.Parsers.DocxResumeParser import DocxResumeParser
from MatchCV.Parser.Parsers.PdfResumeParser import PdfResumeParser
from MatchCV.Parser.Services.ResumeParserService import ResumeParserService


VALID_TEXT = (
    "Experiencia profissional em desenvolvimento de sistemas "
    "com Python, APIs REST, bancos de dados e testes automatizados."
)


def create_pdf_bytes(text: str = VALID_TEXT) -> bytes:
    """Gera um PDF real e mínimo, com texto extraível, sem fixture binária."""
    escaped_text = (
        text.replace("\\", "\\\\")
        .replace("(", "\\(")
        .replace(")", "\\)")
    )
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped_text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
        + stream
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    pdf = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode())
        pdf.extend(obj)
        pdf.extend(b"\nendobj\n")

    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode())
    pdf.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode()
    )
    return bytes(pdf)


def create_docx_bytes(text: str = VALID_TEXT) -> bytes:
    document = Document()
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def create_encrypted_pdf_bytes() -> bytes:
    reader = PdfReader(BytesIO(create_pdf_bytes()))
    writer = PdfWriter()
    writer.clone_document_from_reader(reader)
    writer.encrypt("senha-de-teste")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def create_service() -> ResumeParserService:
    return ResumeParserService(
        pdf_parser=PdfResumeParser(),
        doc_parser=DocResumeParser(),
        docx_parser=DocxResumeParser(),
    )


class TestRealResumeFiles:
    def test_extracts_text_from_real_pdf(self):
        content = create_pdf_bytes()
        result = create_service().parse(
            file_stream=BytesIO(content),
            file_name="curriculo.pdf",
            file_size_bytes=len(content),
        )

        assert result.file_type.value.lower() == "pdf"
        assert result.is_text_extractable is True
        assert VALID_TEXT in result.extracted_text
        assert result.page_count == 1

    def test_extracts_text_from_real_docx(self):
        content = create_docx_bytes()
        result = create_service().parse(
            file_stream=BytesIO(content),
            file_name="curriculo.docx",
            file_size_bytes=len(content),
        )

        assert result.file_type.value.lower() == "docx"
        assert result.is_text_extractable is True
        assert VALID_TEXT in result.extracted_text

    @pytest.mark.parametrize(
        ("file_name", "content", "expected_message"),
        [
            (
                "corrompido.pdf",
                b"isto nao e um PDF",
                "Não foi possível ler o arquivo PDF",
            ),
            (
                "corrompido.docx",
                b"isto nao e um DOCX",
                "Não foi possível ler o arquivo DOCX",
            ),
        ],
    )
    def test_rejects_corrupted_files_with_controlled_error(
        self, file_name, content, expected_message
    ):
        with pytest.raises(ParserException, match=expected_message):
            create_service().parse(
                file_stream=BytesIO(content),
                file_name=file_name,
                file_size_bytes=len(content),
            )

    def test_rejects_password_protected_pdf(self):
        content = create_encrypted_pdf_bytes()

        with pytest.raises(ParserException, match="protegido por senha"):
            create_service().parse(
                file_stream=BytesIO(content),
                file_name="protegido.pdf",
                file_size_bytes=len(content),
            )

    def test_rejects_empty_docx_after_successful_file_read(self):
        document = Document()
        stream = BytesIO()
        document.save(stream)
        content = stream.getvalue()

        with pytest.raises(ParserException, match="Não foi possível extrair texto"):
            create_service().parse(
                file_stream=BytesIO(content),
                file_name="vazio.docx",
                file_size_bytes=len(content),
            )

    def test_rejects_text_below_minimum_useful_character_count(self):
        content = create_pdf_bytes("Texto curto 123")

        with pytest.raises(ParserException, match="30 caracteres úteis"):
            create_service().parse(
                file_stream=BytesIO(content),
                file_name="texto-insuficiente.pdf",
                file_size_bytes=len(content),
            )

    def test_rejects_file_above_one_mb(self):
        content = b"x" * (1_048_576 + 1)

        with pytest.raises(ParserException, match="ultrapassar 1 MB"):
            create_service().parse(
                file_stream=BytesIO(content),
                file_name="curriculo.pdf",
                file_size_bytes=len(content),
            )

    @pytest.mark.parametrize(
        ("file_name", "content", "expected_message"),
        [
            (
                "arquivo.docx",
                create_pdf_bytes(),
                "Não foi possível ler o arquivo DOCX",
            ),
            (
                "arquivo.pdf",
                create_docx_bytes(),
                "Não foi possível ler o arquivo PDF",
            ),
        ],
    )
    def test_rejects_content_whose_real_format_does_not_match_extension(
        self, file_name, content, expected_message
    ):
        with pytest.raises(ParserException, match=expected_message):
            create_service().parse(
                file_stream=BytesIO(content),
                file_name=file_name,
                file_size_bytes=len(content),
            )

    def test_confirms_legacy_doc_is_not_supported_yet(self):
        content = b"conteudo de exemplo de um arquivo DOC legado"

        with pytest.raises(
            ParserException,
            match="Arquivos DOC legados não podem ser processados",
        ):
            create_service().parse(
                file_stream=BytesIO(content),
                file_name="curriculo.doc",
                file_size_bytes=len(content),
            )
